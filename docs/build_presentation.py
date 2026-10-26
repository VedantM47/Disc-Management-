# Builds docs/SpaceRescue_Presentation.pptx. Run: python docs/build_presentation.py
# All numbers come from PROGRESS.md / STATUS_REPORT.md or measurements logged there; unmeasured values say TBD.
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SpaceRescue_Presentation.pptx")
SCREENS = os.path.join(HERE, "screens")

# --- Visual system -------------------------------------------------------------------
NAVY, LIGHT, WHITE, BORDER = "0F1B2D", "EEF1F5", "FFFFFF", "D5DCE5"
ORANGE, RED, TEAL, BLUE = "F5A623", "E5484D", "1F9E89", "3E6FD8"
MUTED, NAVY2 = "5B6B80", "1B2B44"
TITLE_FONT, BODY_FONT, ICON_FONT = "Cambria", "Calibri", "Segoe MDL2 Assets"
TITLE_PT, HEAD_PT, BIG_PT, BODY_PT, SMALL_PT, FOOT_PT = 36, 22, 44, 16, 14, 10
FOOTER = "SpaceRescue · OS Course Project"
W, H = 10.0, 5.625


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
slide_no = [0]


def text(slide, x, y, w, h, parts, size=BODY_PT, color=NAVY, font=BODY_FONT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """parts: str, or list of paragraphs; a paragraph is str or list of (text, {opts}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = parts if isinstance(parts, list) else [parts]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    return s


def card(slide, x, y, w, h, fill=WHITE):
    return rect(slide, x, y, w, h, fill, BORDER if fill == WHITE else None)


def circle(slide, x, y, d, fill, label="", size=14, color=WHITE, font=BODY_FONT, bold=True):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = rgb(color)
    return s


def icon(slide, x, y, glyph, fill=BLUE, d=0.5):
    return circle(slide, x, y, d, fill, glyph, size=int(d * 34), font=ICON_FONT, bold=False)


def badge(slide, x, y, n, d=0.42):
    return circle(slide, x, y, d, ORANGE, f"{n:02d}", size=14, color=NAVY)


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.75, arrow=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if arrow:
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    if dash:
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.insert(1, d)
    return c


def base(dark=False, title=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else LIGHT)
    slide_no[0] += 1
    if title:
        text(s, 0.5, 0.3, 9.0, 0.7, title, size=TITLE_PT, color=WHITE if dark else NAVY, font=TITLE_FONT, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    if slide_no[0] > 1:
        text(s, 0.5, 5.3, 5, 0.2, FOOTER, size=FOOT_PT, color="8A9AAF" if dark else MUTED)
        text(s, 8.5, 5.3, 1.0, 0.2, str(slide_no[0]), size=FOOT_PT, color="8A9AAF" if dark else MUTED, align=PP_ALIGN.RIGHT)
    return s


def strip(slide, y, lead, rest, h=0.55):
    rect(slide, 0.5, y, 9.0, h, NAVY)
    text(slide, 0.75, y, 8.5, h, [[(lead + " ", {"bold": True, "color": ORANGE}), (rest, {"color": WHITE})]],
         size=17, anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def grid_cards(slide, items, top=1.25, cols=3, rows=2, card_h=None, gap=0.2, glyph_color=BLUE, head=22, body=16):
    """items: (glyph or number, heading, line, file)."""
    cw = (9.0 - gap * (cols - 1)) / cols
    ch = card_h or ((5.15 - top - gap * (rows - 1)) / rows)
    for i, it in enumerate(items):
        x = 0.5 + (i % cols) * (cw + gap)
        y = top + (i // cols) * (ch + gap)
        card(slide, x, y, cw, ch)
        g, head_t, line_t, file_t = (list(it) + [""] * 4)[:4]
        if isinstance(g, int):
            badge(slide, x + 0.18, y + 0.18, g)
        elif g:
            icon(slide, x + 0.18, y + 0.18, g, glyph_color, d=0.42)
        text(slide, x + 0.72, y + 0.12, cw - 0.85, 0.72, head_t, size=head, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        paras = []
        if line_t:
            paras.append([(line_t, {})])
        if file_t:
            if len(file_t) > 20 and "." in file_t.split()[0]:
                file_t = file_t.split()[0].rstrip(",")  # file name on the slide; functions are in VIVA_QA.md
            paras.append([(file_t, {"size": SMALL_PT, "color": MUTED, "font": "Consolas"})])
        text(slide, x + 0.2, y + 0.92, cw - 0.38, ch - 0.98, paras, size=body, spacing=3)


def table(slide, x, y, w, col_w, rows, header_pt=15, body_pt=15, row_h=0.42, center_from=1, color_marks=False):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No style, table grid
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (WHITE if r % 2 else "F6F8FB"))
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (c >= center_from and r > 0 and color_marks) or (r == 0 and c >= center_from and color_marks) else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = val
            run.font.name = BODY_FONT
            run.font.size = Pt(header_pt if r == 0 else body_pt)
            run.font.bold = r == 0 or (c == 0 and color_marks) or (color_marks and r == len(rows) - 1)
            color = WHITE if r == 0 else NAVY
            if color_marks and r > 0 and c >= center_from:
                color = {"✓": TEAL, "✗": RED, "~": ORANGE}.get(val[:1], NAVY)
                run.font.size = Pt(18)
                run.font.bold = True
            run.font.color.rgb = rgb(color)
    return tbl


def picture(slide, name, x, y, w):
    path = os.path.join(SCREENS, name)
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = rgb(BORDER)
    pic.line.width = Pt(1)
    return pic


def ring(slide, cx, cy, d, used_fraction, used_color, label, sub):
    """Drive ring motif: a donut whose arc shows the used share."""
    full = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    full.fill.solid(); full.fill.fore_color.rgb = rgb("2A3B55"); full.line.fill.background(); full.adjustments[0] = 0.13
    arc = slide.shapes.add_shape(MSO_SHAPE.BLOCK_ARC, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    arc.fill.solid(); arc.fill.fore_color.rgb = rgb(used_color); arc.line.fill.background()
    arc.adjustments[0] = 270.0
    arc.adjustments[1] = (270.0 + 360.0 * used_fraction) % 360.0
    arc.adjustments[2] = 0.13
    text(slide, cx - d / 2, cy - 0.32, d, 0.4, label, size=22, color=WHITE, font=TITLE_FONT, bold=True, align=PP_ALIGN.CENTER)
    text(slide, cx - d / 2, cy + 0.08, d, 0.3, sub, size=14, color="B8C4D6", align=PP_ALIGN.CENTER)


# =====================================================================================
# 1. Title
s = base(dark=True)
text(s, 0.6, 0.75, 5, 0.3, "OPERATING SYSTEMS COURSE PROJECT", size=13, color=ORANGE, bold=True)
text(s, 0.6, 1.15, 5.6, 0.9, "SpaceRescue", size=54, color=WHITE, font=TITLE_FONT, bold=True)
text(s, 0.6, 2.1, 5.2, 1.0, "Disk management that still works when the disk is full", size=22, color="DCE3EE")
text(s, 0.6, 3.75, 5.5, 0.35, "Vedant  ·  Shreeviraj", size=18, color=WHITE, bold=True)
text(s, 0.6, 4.15, 5.5, 0.3, "October 2026", size=14, color="8A9AAF")
ring(s, 7.6, 2.55, 2.5, 0.998, RED, "C: 512 MB", "free of 292 GB")
rect(s, 6.55, 4.15, 2.1, 0.42, RED, radius=0.21)
text(s, 6.55, 4.15, 2.1, 0.42, "Disk is full", size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "SpaceRescue is a Windows disk manager built for the moment the system drive is full. "
         "At that moment Windows updates fail, apps crash, and most disk tools also fail because they need disk space and memory themselves. "
         "Our tool scans fast, explains what is safe to remove, and writes nothing to the full drive. "
         "It also shows the operating-system internals behind disk space: clusters, the MFT, extents and slack.")

# 2. Agenda
s = base(title="Agenda")
agenda = ["Introduction", "Problem statement", "Gaps in Windows tools", "Use cases and objectives", "Related work",
          "OS concepts in our code", "Built for 0 bytes free", "Algorithms", "Methodology and architecture",
          "The product", "Results", "Conclusion"]
for i, a in enumerate(agenda):
    col, row = i // 6, i % 6
    x, y = 0.5 + col * 4.6, 1.2 + row * 0.66
    card(s, x, y, 4.4, 0.54)
    badge(s, x + 0.1, y + 0.06, i + 1)
    text(s, x + 0.65, y, 3.6, 0.54, a, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "We start with the problem and why Windows' own tools do not solve it, then the users we designed for. "
         "The middle of the talk is the operating-systems content: file systems, concurrency, memory and scheduling, each tied to our code. "
         "Then algorithms, architecture, the product itself, measured results and an honest conclusion.")

# 3. Introduction
s = base(title="Introduction")
card(s, 0.5, 1.2, 4.35, 2.85)
icon(s, 0.75, 1.42, "", BLUE)
text(s, 1.4, 1.42, 3.3, 0.5, "What the OS sees", size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 2.15, 3.85, 1.8, ["Clusters, extents and a free-cluster count", "An MFT record for every file",
                                "Exact bytes allocated on disk"], size=18, spacing=8)
card(s, 5.15, 1.2, 4.35, 2.85)
icon(s, 5.4, 1.42, "", RED)
text(s, 6.05, 1.42, 3.3, 0.5, "What the user sees", size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 5.45, 2.15, 3.85, 1.8, ["“Low disk space”", "No idea which folders grew",
                                 "No idea what is safe to delete"], size=18, spacing=8)
strip(s, 4.3, "SpaceRescue connects the two:", "it reads what the OS knows and explains it in plain words.")
notes(s, "Underneath, Windows knows exactly where every byte is: NTFS keeps a record per file in the Master File Table, "
         "and the volume tracks free clusters. The user only gets a red bar in Explorer. "
         "SpaceRescue reads the same structures the OS uses and turns them into answers a person can act on safely.")

# 4. Problem statement
s = base(title="Problem statement")
card(s, 0.5, 1.15, 9.0, 0.75)
icon(s, 0.68, 1.27, "", RED)
text(s, 1.3, 1.15, 8.0, 0.75, "C: has 512 MB free. A Windows update fails, apps start crashing, the user opens a disk tool…",
     size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
pains = [("", "The tool needs space too", "Temp files, logs and RAM. The pagefile can't grow on a full disk."),
         ("", "Scanning takes minutes", "Our first prototype: 289 s for C: with a cold cache."),
         ("", "Nothing says what is safe", "Users delete the wrong files or give up.")]
for i, (g, h_, b) in enumerate(pains):
    x = 0.5 + i * 3.07
    card(s, x, 2.1, 2.86, 2.0)
    icon(s, x + 0.18, 2.28, g, RED, 0.45)
    text(s, x + 0.75, 2.24, 2.0, 0.55, h_, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 2.9, 2.5, 1.15, b, size=16)
strip(s, 4.3, "Impact:", "failed updates, crashing apps, lost work.")
notes(s, "This is the scenario we designed around. With almost no free space, Windows itself becomes unstable. "
         "Disk tools make it worse: they write temp files, caches and settings, and they need memory exactly when the pagefile cannot grow. "
         "Our own first prototype needed 289 seconds to scan C: from a cold cache, measured on the test laptop. "
         "And no tool tells the user what is safe to remove, so people delete the wrong things.")

# 5. Windows tools gap table
s = base(title="Gaps in Windows’ built-in tools")
Y, N, P = "✓", "✗", "~"
rows = [["Feature", "Explorer\nProperties", "Settings\nStorage", "Disk\nCleanup", "Disk\nMgmt", "Optimize\nDrives", "Space\nRescue"],
        ["Shows where space went", P, P, N, N, N, Y],
        ["Visual map of the disk", N, N, N, N, N, Y],
        ["Works at 0 bytes free", Y, Y, Y, Y, P, Y + "*"],
        ["Safe review + undo", P, P, P, N, N, Y],
        ["Finds app and dev caches", N, P, N, N, N, Y],
        ["Explains system files", N, P, P, N, N, Y],
        ["Early low-space alerts", N, P, N, N, N, Y],
        ["Shows disk internals", P, N, N, P, P, Y]]
table(s, 0.5, 1.1, 9.0, [2.8, 1.03, 1.03, 1.03, 1.03, 1.05, 1.03], rows, header_pt=14, body_pt=16, row_h=0.39, color_marks=True)
text(s, 0.5, 4.72, 9.0, 0.3, "✓ yes   ~ partly   ✗ no      * zero-write design; 0-byte VHDX test pending",
     size=SMALL_PT, color=MUTED)
notes(s, "Rows are features, columns are the Windows 11 tools a user would try. Explorer's folder Properties gives one folder's size at a time. "
         "Settings > Storage shows categories and can clean temporary files; Storage Sense can run when space is low, so alerts are 'partly'. "
         "Disk Cleanup removes a fixed list of categories without undo. Disk Management shows partitions and Optimize Drives shows fragmentation percentage, "
         "which is why they get 'partly' for internals. The gap we close: none of them shows where space went at file level with a safe, undoable cleanup "
         "and warnings before the disk is full. Our 'works at 0 bytes' is a design property; the 0-byte VHDX test is still pending.")

# 6-7. Use cases
for page, cases in ((1, [("", "Student, 512 MB free", "Rescue Mode releases the reserve file; quick wins show in 3 s."),
                         ("", "Developer", "node_modules, Gradle, pip and WSL disks grouped by project and last opened."),
                         ("", "Gamer", "Steam shader cache, unfinished downloads and Epic web cache, each reviewable."),
                         ("", "“pagefile.sys is 7.8 GB?”", "Lock badge, plain explanation, and the safe Windows setting.")]),
                    (2, [("", "Disk filling for weeks", "Quiet alerts at 50, 20, 10, 5, 3, 2, 1 GB and 512 MB."),
                         ("", "Deleted the wrong thing", "Undo from the toast; quarantine keeps a verified copy on another drive."),
                         ("", "Second drive or USB stick", "Work volume holds quarantine and logs; portable mode runs from USB."),
                         ("", "Right-click a folder", "“Analyze with SpaceRescue” scans just that folder (built, untested on a clean PC).")])):
    s = base(title=f"Use cases ({page}/2)")
    for i, (g, h_, b) in enumerate(cases):
        x, y = 0.5 + (i % 2) * 4.6, 1.2 + (i // 2) * 1.95
        card(s, x, y, 4.4, 1.75)
        icon(s, x + 0.2, y + 0.22, g, TEAL)
        text(s, x + 0.85, y + 0.2, 3.4, 0.55, h_, size=21, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, y + 0.85, 3.95, 0.85, b, size=17)
    notes(s, "Every scenario on this slide is supported by the current build. "
             + ("The student case is the core one: Rescue Mode turns on below 1 GB or 2% free, and the quick-win pre-scan answers in about three seconds. "
                "Developer caches are found by a rule table plus folder rules like node_modules next to package.json. "
                "System files such as pagefile.sys get a lock and the safe Windows tool instead of a delete button."
                if page == 1 else
                "Alerts run in a 1.6 MB tray process with hysteresis so they do not repeat. "
                "Quarantine copies to another drive, checks the size, then deletes, and can be undone. "
                "The Explorer right-click entry is implemented in setup.cpp but we have not yet tested it on a clean machine, so we say so."))

# 8. User journey
s = base(title="Thinking from the user’s side")
steps = [("", "Notice", "Tray alert before the disk is full", "Alerts only on the way down"),
         ("", "Understand", "Where space went, in seconds", "Answers stream during the scan"),
         ("", "Act safely", "Review list, quarantine, undo", "Delete is never one click away"),
         ("", "Stay safe", "Reserve file, next alert armed", "Nothing written to the full drive")]
for i, (g, h_, what, why) in enumerate(steps):
    x = 0.5 + i * 2.32
    card(s, x, 1.25, 2.1, 2.45)
    icon(s, x + 0.78, 1.42, g, BLUE, 0.55)
    text(s, x + 0.1, 2.05, 1.9, 0.45, h_, size=22, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 2.55, 1.8, 1.1, what, size=16, align=PP_ALIGN.CENTER)
    if i < 3:
        line(s, x + 2.12, 2.47, x + 2.3, 2.47, ORANGE, 2.5)
    rect(s, x, 3.85, 2.1, 0.95, NAVY2)
    text(s, x + 0.1, 3.85, 1.9, 0.95, why, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.5, 4.88, 9, 0.3, "Top: what the user does.  Bottom: the design rule it forced.", size=SMALL_PT, color=MUTED)
notes(s, "We designed the flow from the user's side. First they notice, through a quiet tray alert, before the disk is full. "
         "Then they understand where space went; that is why results stream while the scan runs. "
         "Then they act, but only through a review list with confirmation and undo, never a one-click delete from the dashboard. "
         "Finally they stay safe: the reserve file can be re-created and the next alert re-arms only after space recovers.")

# 8b. Our solution: each problem mapped to how SpaceRescue answers it
s = base(title="Our solution")
sol = [("", "The tool needs space and RAM", "", "Zero-write Rescue Mode + 512 MB reserve file"),
       ("", "Scanning takes minutes", "", "8-thread scan of C: in 8.2 s; first answer in 3 s"),
       ("", "Nothing says what is safe", "", "Safety badges, review list, quarantine and undo"),
       ("", "System files look scary", "", "Lock badge + safe Windows tool, never Delete"),
       ("", "No warning until it’s full", "", "Tray alerts at 50, 20, 10, 5, 3, 2, 1 GB, 512 MB")]
text(s, 0.5, 1.02, 3.6, 0.3, "PROBLEM", size=13, color=RED, bold=True)
text(s, 4.55, 1.02, 4.9, 0.3, "SPACERESCUE", size=13, color=TEAL, bold=True)
for i, (g1, prob, g2, ans) in enumerate(sol):
    y = 1.32 + i * 0.6
    card(s, 0.5, y, 3.6, 0.5)
    icon(s, 0.58, y + 0.06, g1, RED, 0.38)
    text(s, 1.05, y, 3.0, 0.5, prob, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    line(s, 4.13, y + 0.25, 4.5, y + 0.25, ORANGE, 2.5)
    rect(s, 4.55, y, 4.95, 0.5, NAVY)
    icon(s, 4.63, y + 0.06, g2, TEAL, 0.38)
    text(s, 5.1, y, 4.3, 0.5, ans, size=16, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
strip(s, 4.4, "One tool, one flow:", "notice → understand → act safely → stay safe.")
notes(s, "This is our answer to every problem we raised. Because the tool itself needs space and memory, Rescue Mode writes nothing to the full drive "
         "and can hand back a 512 MB reserve file instantly. Because scanning took minutes, the scanner now uses eight threads and finishes C: in 8.2 seconds, "
         "with a first answer after three seconds. Because users don't know what is safe, every item carries a safety badge and every action can be undone. "
         "System files get a lock and the Windows tool that shrinks them safely, and tray alerts warn long before the disk is full.")

# 9. Objectives
s = base(title="Objectives")
objs = [(1, "Rescue", "Work at 0 bytes free: zero writes to the full drive"),
        (2, "Fast scan", "Useful answer in seconds; full C: well under a minute"),
        (3, "Explain", "Say what each big item is and whether it is safe"),
        (4, "Clean safely", "Review, quarantine, compress; undo every step"),
        (5, "Warn early", "Alerts at fixed thresholds with hysteresis"),
        (6, "Teach internals", "Clusters, slack, extents and MFT on your own disk")]
grid_cards(s, objs, top=1.2, head=22, body=17)
notes(s, "Six objectives drove the work. Rescue and fast scan are the hard engineering ones. "
         "Explain and clean safely are about trust. Warn early prevents the problem, and teach internals is the operating-systems angle: "
         "the same data we use to free space also shows how the file system allocates it.")

# 10. Related work
s = base(title="Related work")
rows = [["Tool we studied", "Does well", "Gap that remains"],
        ["WinDirStat", "Tree list + file types + treemap", "No alerts, no rescue mode"],
        ["WizTree", "Very fast scans via the MFT", "No safety guidance on what to delete"],
        ["TreeSize", "Detailed folder reports and exports", "No early warning before disk is full"],
        ["SpaceSniffer", "Zoomable live treemap", "Shows sizes, not what is safe"],
        ["ncdu (Linux)", "Tiny, fast terminal browser", "Linux only; no Windows system-file awareness"]]
table(s, 0.5, 1.15, 9.0, [2.2, 3.4, 3.4], rows, header_pt=16, body_pt=16, row_h=0.5)
strip(s, 4.35, "Common gap:", "none warn before the disk fills or work around a full system drive.")
notes(s, "We studied five well-known tools. WinDirStat's three linked views shaped our Explore tab, and its cushion treemap paper trail led us to the "
         "van Wijk and Bruls papers. WizTree shows how fast MFT reading can be. None of them warns before the disk is full, "
         "and none is designed to run when the system drive itself has no free space. The gap column is our reading of their feature sets, "
         "not a benchmark of them.")

# 11-13. OS concepts: storage
storage = [
    [("", "MFT ≈ inode table", "One record per file; we read it directly", "scanner_mft.cpp ScanMft"),
     ("", "Directory entries", "Name, size, times, attributes per entry", "scanner.cpp Enumerate"),
     ("", "Extents vs fragments", "Contiguous runs; non-contiguous = fragmented", "internals.cpp QueryFile"),
     ("", "Free-space management", "Free-cluster count × cluster size", "internals.cpp QueryVolume"),
     ("", "Internal fragmentation", "Slack: unused tail of the last cluster", "internals.cpp EstimateSlack"),
     ("", "Reserved blocks", "Reserve file ≈ ext4 blocks reserved for root", "rescue.cpp CreateReserve")],
    [("", "Page cache", "Same scan: 289 s cold vs 58 s warm", "PROGRESS.md baseline"),
     ("", "Hard links", "MFT counts a file once; names count twice", "scanner_mft.cpp"),
     ("", "Reparse points", "Junctions not followed: no cycles", "scanner.cpp NF_REPARSE"),
     ("", "Compression, sparse", "Allocated size can be < logical size", "actions.cpp Compress"),
     ("", "Logical vs on disk", "Node.size vs Node.alloc (AllocationSize)", "scanner.h struct Node"),
     ("", "Journaled rename", "Recycle = same-volume rename, no new space", "actions.cpp recycle undo")],
    [("", "Seek time, scheduling", "One sequential MFT read beats 400K random reads", "scanner_mft.cpp"),
     ("", "Unbuffered I/O", "FILE_FLAG_NO_BUFFERING, cluster-aligned 4 MB", "scanner_mft.cpp"),
     ("", "Overlapped I/O", "8 reads in flight; OVERLAPPED + event", "scanner_mft.cpp workers"),
     ("", "Restore points", "Shadow copies listed with the safe tool", "cleanup.cpp AnalyzeSystem")]]
for k, items in enumerate(storage):
    s = base(title=f"OS concepts: file systems ({k + 1}/3)")
    grid_cards(s, items, top=1.15, rows=2 if len(items) > 4 else 2, cols=3 if len(items) > 4 else 2)
    notes(s, ["The NTFS Master File Table is the counterpart of the Unix inode table: one record per file with its attributes and where its data lives. "
              "Directory enumeration returns entries with name, size, timestamps and attributes. "
              "File data is stored as extents, runs of consecutive clusters; a file split into many extents is fragmented, and Disk Internals draws that strip. "
              "Free space is managed in clusters, and slack is internal fragmentation. Our reserve file is the same idea as ext4 reserving blocks for root.",
              "The page cache is why the same scan took 289 seconds cold and 58 warm. Hard links are one file with several names: "
              "enumeration counts each name, the MFT counts the file once, which explained a 15 GB difference on C:. "
              "Junctions are not followed so we never loop or double count. Compressed and sparse files allocate less than their logical size. "
              "Recycling on the same volume is a rename, a journaled metadata operation, so it needs no new space; undo renames the file back.",
              "Disk scheduling: random directory reads cause many seeks, a sequential MFT read does not; FCFS, SSTF and SCAN from the syllabus all try to reduce "
              "that seek cost. The MFT reader uses unbuffered, cluster-aligned reads and overlapped I/O with eight reads in flight. "
              "The System tab lists restore points with the Windows tool to shrink them."][k])

# 14-16. Concurrency
s = base(title="Concurrency: the parallel scanner")
# workers -> queue -> tree diagram
for i in range(3):
    y = 1.25 + i * 0.95
    rect(s, 0.5, y, 1.9, 0.75, BLUE)
    text(s, 0.5, y, 1.9, 0.75, f"Worker {i + 1}" if i < 2 else "Worker 8", size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    line(s, 2.4, y + 0.38, 3.3, 2.2, MUTED, 1.5)
    line(s, 6.7, 2.2, 7.45, y + 0.38, MUTED, 1.5, dash=True)
text(s, 0.5, 4.15, 1.9, 0.3, "… min(cores, 8)", size=14, color=MUTED, align=PP_ALIGN.CENTER)
rect(s, 3.3, 1.45, 3.4, 1.55, NAVY)
text(s, 3.45, 1.52, 3.1, 0.45, "Folder queue (LIFO)", size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, 3.45, 1.97, 3.1, 0.95, ["SRW lock + condition variable", "push sub-folders, pop next"], size=16, color="DCE3EE", align=PP_ALIGN.CENTER)
rect(s, 7.45, 1.25, 2.05, 2.65, TEAL)
text(s, 7.5, 1.35, 1.95, 2.5, ["Scan tree", ("• atomic fetch_add reserves node blocks"),
                               "• per-worker name arenas"], size=15, color=WHITE, spacing=6)
rect(s, 3.3, 3.25, 3.4, 0.9, ORANGE)
text(s, 3.3, 3.25, 3.4, 0.9, ["UI thread", "PostMessage when done"], size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
line(s, 5.0, 3.0, 5.0, 3.25, ORANGE, 2)
strip(s, 4.5, "No lock is held during disk I/O:", "workers enumerate a folder first, then lock only to push results.")
notes(s, "This is scanner.cpp. Run starts min(cores, 8) worker threads. They share one LIFO queue of folders protected by a std::mutex and a condition variable; "
         "in our build these compile to a Windows SRW lock and SleepConditionVariableSRW, which we confirmed in the executable's import table. "
         "A worker pops a folder, enumerates it with no lock held, reserves a contiguous block of nodes with one atomic fetch_add, and writes names into its own arena. "
         "Then it takes the lock briefly to push the sub-folders. When the queue is empty and no worker is busy, the condition variable wakes everyone to exit, "
         "and the scanner posts a message to the UI thread.")

conc = [
    [("", "Thread pool", "min(cores, 8) std::thread workers", "scanner.cpp Run"),
     ("", "Producer–consumer", "Workers push sub-folders, pop folders", "scanner.cpp m_queue"),
     ("", "Mutual exclusion", "std::mutex = SRW lock (verified import)", "scanner.h m_qLock"),
     ("", "Condition variable", "Idle and termination: queue empty, busy == 0", "scanner.cpp m_cv.wait"),
     ("", "Atomic operations", "fetch_add for progress and node blocks", "scanner.cpp ReserveNodes"),
     ("", "Thread-local storage", "64 KiB enumeration buffer per worker", "scanner.cpp t_scratch")],
    [("", "Races prevented", "Each worker owns its folder’s child list", "scanner.cpp ScanFolder"),
     ("", "Deadlock avoided", "No nested locks; no lock during I/O", "scanner.cpp ScanFolder"),
     ("", "Fairness, load balance", "Shared LIFO queue: idle workers take next folder", "scanner.cpp worker loop"),
     ("", "Thread-safe UI", "Workers post messages; UI owns windows", "preview.cpp"),
     ("", "Cancellation", "Atomic flag checked between folders", "scanner.cpp m_cancel"),
     ("", "Amdahl, contention", "8 threads 8.2 s; 12 threads 11.8 s", "warm C:, measured")]]
for k, items in enumerate(conc):
    s = base(title=["Concurrency: primitives", "Concurrency: correctness"][k])
    grid_cards(s, items, top=1.15)
    notes(s, ["The worker pool, the producer-consumer queue, mutual exclusion, the condition variable, atomics and thread-local storage are all in scanner.cpp. "
              "We checked the compiled executable: it imports AcquireSRWLockExclusive and SleepConditionVariableSRW, so the C++ mutex and condition variable "
              "are Windows slim reader-writer locks used in exclusive mode. Progress counters and node reservation use atomic fetch_add, so the hot path has no lock.",
              "Races: two workers never touch the same folder's child list, because only the worker scanning a folder links its children. "
              "Deadlock: there is never more than one lock held at a time, and none is held during I/O. The shared queue balances load automatically. "
              "Workers never touch windows; they post messages. Amdahl's law shows up in our numbers: beyond 8 threads the disk and the queue lock become the "
              "serial part, and 12 threads were slower than 8. Files in use are opened with full sharing flags, and access-denied folders are counted, not fatal."][k])

# 17-18. Processes, memory, scheduling, protection
proc = [
    [("", "I/O and CPU priority", "Background mode lowers CPU and I/O priority", "scanner.cpp, quickwins.cpp"),
     ("", "Context switches", "More threads than cores costs time: 12 > 8", "measured on C:"),
     ("", "Virtual memory", "Full disk: the pagefile can’t grow, less RAM", "why we cap memory"),
     ("", "Working set", "Peak tracked; tray trims to 1.6 MB private", "common.cpp, tray.cpp"),
     ("", "Thrashing avoided", "Memory budget folds small files into one node", "scanner.cpp ScanFolder"),
     ("", "Arena allocation", "1 MB chunks, 64-byte nodes, no per-file malloc", "scanner.h ScanTree")],
    [("", "System calls", "NtQueryDirectoryFile and NTFS FSCTLs", "kernel mode ↔ user mode"),
     ("", "Hiberfil and swapfile", "Shown with the safe Windows setting", "cleanup.cpp AnalyzeSystem"),
     ("", "Privilege and UAC", "Raw volume read needs admin; else fallback", "scanner.cpp, common.cpp"),
     ("", "Access control", "520 folders access-denied: counted, not fatal", "scanner.cpp NF_ERROR"),
     ("", "Protected files", "Delete disabled for Windows-managed paths", "cleanup.cpp IsProtectedPath"),
     ("", "Message loop, timers", "WM_APP messages; 30 s SetTimer poll in tray", "tray.cpp, preview.cpp")]]
for k, items in enumerate(proc):
    s = base(title=["Memory and scheduling", "Protection and system calls"][k])
    grid_cards(s, items, top=1.15)
    notes(s, ["In Rescue Mode the scan threads enter background mode, which lowers both CPU and I/O priority so the PC stays usable. "
              "Thread count is min(cores, 8) because extra threads only add context switches; we measured 12 threads slower than 8. "
              "A full disk also means less memory, because the pagefile cannot grow, so we cap the tree with a memory budget and fold small files. "
              "Nodes and names live in 1 MB arenas instead of millions of small heap allocations, which also avoids heap fragmentation.",
              "Every scan crosses into kernel mode through NtQueryDirectoryFile or DeviceIoControl with NTFS FSCTLs. "
              "Reading the raw volume needs administrator rights; without them we fall back to directory enumeration. "
              "Access-denied folders are counted, 520 on our C:, and the scan continues. Windows-managed paths cannot be deleted from the app. "
              "The tray watcher is a message loop with a 30-second timer and Shell_NotifyIcon notifications."][k])

# 19. Built for 0 bytes free
s = base(title="Built for 0 bytes free")
eff = [("", "Zero-write scan", "Tree only in RAM"),
       ("", "64-byte nodes", "Arenas, no per-file malloc"),
       ("", "Memory budget", "Fold “[N small files]”"),
       ("", "No path strings", "Built only when needed"),
       ("", "Streaming", "Sizes appear during scan"),
       ("", "Reserve file", "512 MB back on demand"),
       ("", "Work volume", "Writes go to another drive"),
       ("", "One static exe", "No runtimes or DLLs")]
for i, (g, h_, b) in enumerate(eff):
    x, y = 0.5 + (i % 2) * 4.6, 1.12 + (i // 2) * 0.68
    card(s, x, y, 4.4, 0.58)
    icon(s, x + 0.1, y + 0.08, g, TEAL, 0.42)
    text(s, x + 0.65, y, 1.75, 0.58, h_, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 2.4, y, 1.95, 0.58, b, size=16, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
nums = [("255 MB", "peak RAM, full C: scan"), ("1.6 MB", "tray alert process"), ("19 bytes", "written per scan (0 in Rescue)")]
for i, (n, l) in enumerate(nums):
    x = 0.5 + i * 3.07
    rect(s, x, 3.95, 2.86, 1.12, NAVY)
    text(s, x, 3.98, 2.86, 0.6, n, size=30, color=ORANGE, font=TITLE_FONT, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 4.6, 2.86, 0.4, l, size=15, color=WHITE, align=PP_ALIGN.CENTER)
notes(s, "On a full disk we lose both disk space and usable memory, so every design choice here saves one or the other. "
         "The scan tree never touches disk. Nodes are 64 bytes in 1 MB chunks, ASCII names take one byte per character, and a memory budget folds small files. "
         "Measured on our laptop: 255 MB peak working set for the scanner on 1.68 million files with 8 threads, 1.6 MB for the tray process, "
         "and 19 bytes written during a whole app scan of C: - that is the last-scan-path setting, which is suppressed in Rescue Mode by design.")

# 20-21. Algorithms
algs = [
    [(1, "Parallel tree traversal", "DFS over a shared LIFO queue; O(n)", "scanner.cpp"),
     (2, "MFT parsing", "Fixups, then attributes per record; O(records)", "scanner_mft.cpp"),
     (3, "Size roll-up", "Post-order pass, children first; O(n)", "scanner.cpp RollUp"),
     (4, "Squarified treemap", "Rows by worst aspect ratio; cushion shading", "viz.cpp Treemap"),
     (5, "Sunburst partition", "Angle ∝ size; tiny slices grouped", "viz.cpp Sunburst"),
     (6, "Duplicate funnel", "Size → 64 KB hash → full hash", "insights.cpp")],
    [(7, "Stale-file classifier", "Last access, or last write if updates are off", "insights.cpp"),
     (8, "Extent merging", "Join runs where next LCN = end of previous", "internals.cpp QueryFile"),
     (9, "Top-K largest", "std::partial_sort (heap select); O(n log k)", "internals.cpp"),
     (10, "Alerts with hysteresis", "Fire on crossing down; re-arm at +10%", "tray.cpp AlertLevel")]]
for k, items in enumerate(algs):
    s = base(title=f"Algorithms ({k + 1}/2)")
    grid_cards(s, items, top=1.15, cols=3 if k == 0 else 2)
    notes(s, ["Six algorithms do the heavy lifting. The scan is a depth-first traversal spread over workers. The MFT reader applies the update-sequence fixups "
              "and parses each record's attributes. Sizes are rolled up in one post-order pass. The treemap uses the squarified layout of Bruls, Huizing and van Wijk "
              "with van Wijk's cushion shading; full C: renders in 76 ms. The duplicate finder only reads file contents for files that already share a size.",
              "Stale files use the last-access time only when NTFS updates it; on this PC the registry value 0x80000002 says it does, otherwise we fall back to last write. "
              "Extent merging joins physically adjacent runs; on a 6 GB ISO that turned fsutil's 86 runs into 24 fragments. "
              "Alerts fire once when crossing a threshold downward and re-arm only after free space rises 10% above it; 11 of 11 unit tests pass."][k])

# 22. Deep dive: scan pipeline + thread scaling chart
s = base(title="Deep dive: the scan pipeline")
stages = [("Enumerate", "64 KiB buffer per thread"), ("Reserve", "one atomic node block"),
          ("Push", "sub-folders to the queue"), ("Roll up", "post-order, O(n)")]
for i, (h_, b) in enumerate(stages):
    x = 0.5 + i * 1.18
    rect(s, x, 1.2, 1.05, 1.55, NAVY if i % 2 == 0 else BLUE)
    text(s, x, 1.27, 1.05, 0.4, h_, size=16, color=ORANGE if i % 2 == 0 else WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.06, 1.72, 0.93, 1.0, b, size=14, color=WHITE, align=PP_ALIGN.CENTER)
    if i < 3:
        line(s, x + 1.05, 1.98, x + 1.18, 1.98, ORANGE, 2)
cd = CategoryChartData()
cd.categories = ["1", "2", "4", "8", "12"]
cd.add_series("Seconds", (25.1, 14.4, 11.6, 8.2, 11.8))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.3), Inches(1.1), Inches(4.2), Inches(3.2), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Full C: scan, warm cache (s) by threads"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(NAVY)
plot = ch.plots[0]
plot.has_data_labels = True
plot.data_labels.font.size = Pt(15)
plot.data_labels.font.bold = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.series[0].format.fill.solid()
plot.series[0].format.fill.fore_color.rgb = rgb(BLUE)
pt = plot.series[0].points[3]
pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(ORANGE)
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.category_axis.tick_labels.font.size = Pt(15)
for i, (n, l) in enumerate([("289 s", "first prototype, cold"), ("58 s", "first prototype, warm"), ("8.2 s", "now, 8 threads, warm")]):
    x = 0.5 + i * 1.6
    card(s, x, 3.0, 1.48, 1.3)
    text(s, x, 3.05, 1.48, 0.6, n, size=28, color=ORANGE if i == 2 else NAVY, font=TITLE_FONT, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.08, 3.65, 1.32, 0.6, l, size=14, color=MUTED, align=PP_ALIGN.CENTER)
strip(s, 4.5, "Cold cache with the new scanner:", "TBD. Elevated MFT path: 16.4 s (serial version); parallel TBD.")
notes(s, "Each folder goes through four steps: enumerate with NtQueryDirectoryFile into a 64 KiB per-thread buffer, reserve a contiguous block of nodes "
         "with one atomic operation, push sub-folders, and at the end roll sizes up in post-order. The chart is measured today on the same laptop and drive: "
         "1.68 million files, warm cache. Going from 1 to 8 threads cut 25.1 s to 8.2 s; 12 threads was slower, 11.8 s, because the disk and the queue lock "
         "became the bottleneck. The first prototype took 58 s warm and 289 s cold. We have not yet re-measured the new scanner with a cold cache, "
         "and the parallel MFT reader still needs an elevated benchmark.")

# 23. Methodology
s = base(title="Methodology")
meth = [("Scan", "Parallel enumeration or MFT read"), ("Model", "Compact tree, kinds, roll-up"),
        ("Analyse", "Categories, stale, duplicates"), ("Decide", "Review list with safety badges"),
        ("Act and undo", "Recycle, quarantine, compress")]
for i, (h_, b) in enumerate(meth):
    y = 1.15 + i * 0.78
    card(s, 0.5, y, 5.4, 0.66)
    badge(s, 0.62, y + 0.12, i + 1)
    text(s, 1.2, y, 1.8, 0.66, h_, size=19, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.0, y, 2.8, 0.66, b, size=16, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 6.15, 1.15, 3.35, 3.78, NAVY)
text(s, 6.35, 1.25, 3.0, 0.45, "How we validated", size=20, color=ORANGE, bold=True)
text(s, 6.35, 1.75, 3.0, 3.1, ["Scan benchmarks: threads, cold vs warm", "Extents vs fsutil queryextents",
                                "Alert thresholds: 11/11 tests", "UI click tests with frame snapshots",
                                "Pending: 0-byte VHDX, install"], size=16, color=WHITE, spacing=7)
notes(s, "The pipeline has five stages, from scanning to acting with undo. For validation we benchmarked the scanner at 1 to 12 threads, "
         "compared our extent counts with Windows' fsutil file queryextents, unit-tested the alert thresholds, and drove the interface with posted clicks "
         "while it rendered frames to files. Two tests from our plan are still pending because they need administrator rights: filling a virtual disk "
         "to 0 bytes, and the install-uninstall cycle on a clean PC.")

# 24. Architecture
s = base(title="System architecture")
col = [(0.5, "User", ["Explorer right-click", "Tray alert", "App window"]),
       (2.75, "UI layer", ["Eight tabs (Direct2D)", "Review + undo toast", "Search, Back, Forward"]),
       (5.0, "Engine", ["Parallel scanner + MFT", "Insights and cleanup", "Actions + undo log", "Disk internals", "Alerts"]),
       (7.25, "OS interfaces", ["NtQueryDirectoryFile", "NTFS FSCTLs", "Shell, DWM, Direct2D", "NTFS + work volume"])]
fills = [NAVY2, BLUE, NAVY, TEAL]
for i, (x, head, items) in enumerate(col):
    rect(s, x, 1.15, 2.25, 0.5, fills[i])
    text(s, x, 1.15, 2.25, 0.5, head, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for j, it in enumerate(items):
        y = 1.75 + j * 0.55
        card(s, x, y, 2.25, 0.47)
        text(s, x + 0.08, y, 2.1, 0.47, it, size=15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        line(s, x + 2.27, 1.4, x + 2.48, 1.4, ORANGE, 2.5)
text(s, 5.0, 4.55, 2.25, 0.3, "8 worker threads", size=14, color=MUTED, align=PP_ALIGN.CENTER)
strip(s, 4.55 + 0.0, "Never written to the full drive:", "scan tree, caches, logs. Writes go to the work volume.", h=0.55) if False else None
rect(s, 0.5, 4.6, 4.3, 0.55, NAVY)
text(s, 0.65, 4.6, 4.1, 0.55, [[("Never on the full drive: ", {"bold": True, "color": ORANGE}), ("tree, caches, logs", {"color": WHITE})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Four layers. Users reach the app through the window, the Explorer right-click entry and the tray alert. The UI layer is our own Direct2D interface "
         "with eight tabs, the review list and undo. The engine has the parallel scanner with the MFT fast path, the insights and cleanup rules, "
         "the actions with an undo log, disk internals and alerts. Only the engine talks to the operating system, through directory queries, NTFS FSCTLs, "
         "the Shell and DWM. The scan tree, caches and logs never go to the full drive: quarantine and logs go to the work volume, otherwise they stay in memory.")

# 25-27. Product
s = base(title="The product")
picture(s, "app-overview.png", 0.5, 1.1, 5.9)
text(s, 6.65, 1.2, 2.85, 3.6, ["Sunburst of the drive, free space in the centre",
                               "29 GB safely reclaimable on our C:",
                               "Top space users, click to explore"], size=18, spacing=14)
notes(s, "This is the first screen after a scan of our C:. The sunburst shows each folder level as a ring, coloured by the kind of files inside, with free space in grey. "
         "The meter splits what can be freed into Safe, Review and Caution: 29 GB was safe on this drive. Every card opens the item in Explore.")
s = base(title="Clean up and System")
picture(s, "app-cleanup.png", 0.5, 1.15, 4.4)
picture(s, "app-system.png", 5.1, 1.15, 4.4)
text(s, 0.5, 4.05, 4.4, 0.6, "Four sections with safety badges and per-project groups", size=16)
text(s, 5.1, 4.05, 4.4, 0.6, "Windows uses 70.8 GB; each item gets the safe Windows tool", size=16)
notes(s, "Clean up groups everything removable into four sections, each with a safety badge; clicking a group opens a review list with confirmation and undo. "
         "The System tab shows what Windows itself uses, 70.8 GB here, with a lock on each item and the Windows tool that can shrink it, such as turning off hibernation.")
s = base(title="Disk internals and Rescue Mode")
picture(s, "app-internals.png", 0.5, 1.15, 4.4)
picture(s, "app-rescue.png", 5.1, 1.15, 4.4)
text(s, 0.5, 4.05, 4.4, 0.6, "8 GB system image stored in 58 fragments", size=16)
text(s, 5.1, 4.05, 4.4, 0.6, "Rescue banner: nothing is written to C:", size=16)
notes(s, "Disk Internals shows the cluster size, slack, a cluster map and the real extents of a file: this 8 GB emulator image is stored in 58 pieces. "
         "On the right is Rescue Mode, forced here with a high threshold in a test copy so we could capture it: the banner states that nothing is written to C:, "
         "and the reserve-file button appears when a reserve exists.")

# 28. Tools and technologies
s = base(title="Tools and technologies")
tools = [("", "C++23", "Engine and UI"), ("", "Win32 API", "Files, threads, shell"),
         ("", "Direct2D", "With DirectWrite: GPU-drawn interface"), ("", "clang", "llvm-mingw: one static x64 exe"),
         ("", "PowerShell", "Build, package, tests"), ("", "Git", "GitHub, commits per module"),
         ("", "Claude Code", "Pair programming and code audits"), ("", "Python", "Icon, strings, this deck")]
for i, (g, h_, b) in enumerate(tools):
    x, y = 0.5 + (i % 4) * 2.3, 1.2 + (i // 4) * 1.9
    card(s, x, y, 2.1, 1.7)
    icon(s, x + 0.15, y + 0.18, g, BLUE, 0.48)
    text(s, x + 0.75, y + 0.12, 1.3, 0.6, h_, size=18, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, y + 0.85, 1.85, 0.8, b, size=16, color=MUTED)
notes(s, "Everything is native: C++23 on the Win32 API, with Direct2D and DirectWrite for the interface. "
         "We build with clang from llvm-mingw, installed per user without admin, into one static x64 executable. "
         "PowerShell scripts build, package with Windows' own IExpress, and drive our tests. Claude Code was used as a pair programmer and for code audits.")

# 29. Results
s = base(title="Results")
res = [("58 → 8.2 s", "full C: scan, warm cache"), ("3.0 s", "first “you can free” answer"),
       ("255 MB", "peak RAM, 1.68M files"), ("1.6 MB", "tray alert process"),
       ("76 ms", "treemap of all of C:"), ("TBD", "space freed in 0-byte test")]
for i, (n, l) in enumerate(res):
    x, y = 0.5 + (i % 3) * 3.07, 1.15 + (i // 3) * 1.75
    rect(s, x, y, 2.86, 1.55, NAVY)
    text(s, x, y + 0.15, 2.86, 0.75, n, size=36, color=ORANGE, font=TITLE_FONT, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, y + 0.92, 2.66, 0.55, l, size=16, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0.5, 4.75, 9, 0.35, "Test machine: Intel Core i5-1335U, 16 GB RAM, Windows 11, C: 292 GB NTFS, 1.68M files.",
     size=SMALL_PT, color=MUTED)
notes(s, "All numbers are measured on our test laptop. The full scan of C: went from 58 seconds warm with the first prototype to 8.2 seconds with eight threads. "
         "The quick-win pre-scan showed 12.5 GB safe to free three seconds after launch. The scanner peaked at 255 MB for 1.68 million files, the tray watcher uses 1.6 MB, "
         "and the cushion treemap of the whole drive renders in 76 ms. Space freed in the 0-byte virtual-disk test is still to be measured.")

# 30. Conclusion
s = base(dark=True, title="Conclusion")
text(s, 0.5, 1.15, 4.3, 0.45, "What we delivered", size=22, color=ORANGE, bold=True)
text(s, 0.5, 1.7, 4.3, 3.2, ["Scan of C: in 8.2 s, answer in 3 s", "Zero-write Rescue Mode, reserve file",
                              "Review, quarantine and undo", "Alerts at 8 thresholds", "Disk internals on a real disk"],
     size=17, color=WHITE, spacing=8)
rect(s, 5.1, 1.1, 4.4, 3.85, NAVY2)
text(s, 5.3, 1.2, 4.0, 0.45, "Limitations and future work", size=20, color=ORANGE, bold=True)
text(s, 5.3, 1.75, 4.0, 3.1, ["0-byte VHDX test still to run", "Parallel MFT path: elevated benchmark", "Cold-cache timing of new scanner",
                              "Installer unsigned, per-user only", "Windows 11 top-level menu entry"], size=16, color="DCE3EE", spacing=7)
notes(s, "We delivered a working tool: fast scans, a zero-write rescue mode, safe cleanup with undo, early alerts and an internals view. "
         "We are also clear about what is not done yet. The 0-byte virtual-disk test and the elevated benchmark of the parallel MFT reader need administrator rights, "
         "the cold-cache time of the new scanner is not re-measured, the installer is per-user and unsigned, and the Windows 11 top-level context-menu entry "
         "needs a packaged app identity, which we left for future work.")

# 31. References
s = base(title="References")
refs_l = ["[1] A. Silberschatz, P. B. Galvin, G. Gagne. Operating System Concepts, 10th ed. Wiley, 2018.",
          "[2] B. Shneiderman. Tree visualization with tree-maps: 2-d space-filling approach. ACM Trans. on Graphics 11(1):92–99, 1992.",
          "[3] J. J. van Wijk, H. van de Wetering. Cushion treemaps: visualization of hierarchical information. Proc. IEEE InfoVis, 1999.",
          "[4] M. Bruls, K. Huizing, J. J. van Wijk. Squarified treemaps. Proc. Joint Eurographics/IEEE TCVG VisSym, 2000."]
refs_r = ["[5] J. Stasko, E. Zhang. Focus+context display and navigation techniques for radial, space-filling hierarchy visualizations. Proc. IEEE InfoVis, 2000.",
          "[6] WinDirStat Team. WinDirStat 2.10, windirstat.net (studied as a reference).",
          "[7] Microsoft Learn: FSCTL_GET_NTFS_VOLUME_DATA, FSCTL_GET_RETRIEVAL_POINTERS, NtQueryDirectoryFile.",
          "[8] Microsoft Learn: Slim Reader/Writer (SRW) Locks; Condition Variables; Shell_NotifyIcon."]
text(s, 0.5, 1.15, 4.4, 3.9, refs_l, size=14, spacing=10)
text(s, 5.1, 1.15, 4.4, 3.9, refs_r, size=14, spacing=10)
notes(s, "The OS concepts follow Silberschatz, Galvin and Gagne. The treemap follows Shneiderman's tree-maps, van Wijk and van de Wetering's cushion shading "
         "and the squarified layout of Bruls, Huizing and van Wijk; the sunburst follows Stasko and Zhang. The Windows APIs are documented on Microsoft Learn.")

# 32. Thank you
s = base(dark=True)
text(s, 0.6, 1.5, 5, 1.0, "Thank you", size=54, color=WHITE, font=TITLE_FONT, bold=True)
text(s, 0.6, 2.55, 5, 0.5, "Questions are welcome", size=22, color=ORANGE)
text(s, 0.6, 3.3, 5, 0.4, "SpaceRescue · Operating Systems Course Project", size=15, color="B8C4D6")
ring(s, 7.6, 2.55, 2.5, 0.84, TEAL, "29 GB", "safe to free on C:")
rect(s, 6.55, 4.15, 2.1, 0.42, TEAL, radius=0.21)
text(s, 6.55, 4.15, 2.1, 0.42, "Drive under control", size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Thank you. The ring is the same motif as the title slide, now showing a healthy drive: on our C:, SpaceRescue found 29 GB that is safe to free. We are happy to take questions.")

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
